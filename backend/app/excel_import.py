"""Fillable Excel workbook for a complete account people and organization import."""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from .portable_import import PortableImportBatch


SHEETS = {
    "Account": ("name",),
    "Pods": ("name", "head_stakeholder_id"),
    "Divisions": ("id", "pod", "name", "color", "head_stakeholder_id"),
    "BusinessUnits": ("id", "division_id", "name", "sort_order"),
    "Employees": ("id", "name", "first_name", "last_name", "title", "level", "role", "capability", "location", "active", "manager_employee_id", "skills", "source_record_id"),
    "Stakeholders": ("id", "name", "title", "pod", "division_id", "business_unit_id", "team_type", "organizational_role", "level", "location", "country_code", "manager_stakeholder_id", "is_primary_technology", "is_buyer", "is_influencer", "is_budget_holder", "relationship_strength", "capco_contingents", "budget_amount", "biography", "tags", "source_record_id"),
    "Relationships": ("id", "stakeholder_id", "employee_id", "relationship_role", "is_primary", "effective_from", "effective_to", "is_current"),
}
SHEET_TITLES = {
    "Account": "Account",
    "Pods": "Pods",
    "Divisions": "Divisions",
    "BusinessUnits": "Business Units",
    "Employees": "Employees",
    "Stakeholders": "Client Contacts",
    "Relationships": "Relationships",
}
FRIENDLY_HEADERS = {
    "Account": ("Account name",),
    "Pods": ("Pod name", "Pod head contact ID"),
    "Divisions": ("Division ID", "Pod name", "Division name", "Color (hex)", "Division head contact ID"),
    "BusinessUnits": ("Business unit ID", "Division ID", "Business unit name", "Display order"),
    "Employees": ("Employee ID", "Full name", "First name", "Last name", "Job title", "Level", "Role", "Capability", "Location", "Active?", "Manager employee ID", "Skills (separate with ;)", "Source record ID"),
    "Stakeholders": ("Contact ID", "Full name", "Job title", "Pod name", "Division ID", "Business unit ID", "Team type", "Organization role", "Level", "Location", "Country code", "Manager contact ID", "Primary technology lead?", "Buyer?", "Influencer?", "Budget holder?", "Relationship strength", "Consultants assigned", "Budget amount", "Biography", "Tags (separate with ;)", "Source record ID"),
    "Relationships": ("Relationship ID", "Client contact ID", "Employee ID", "Relationship role", "Primary owner?", "Start date", "End date", "Current?"),
}
REQUIRED_FIELDS = {
    "Account": {"name"},
    "Pods": {"name", "head_stakeholder_id"},
    "Divisions": {"id", "pod", "name"},
    "BusinessUnits": {"id", "division_id", "name"},
    "Employees": {"id", "name", "level", "role", "location"},
    "Stakeholders": {"id", "name", "title", "pod", "team_type", "organizational_role"},
    "Relationships": {"id", "stakeholder_id", "employee_id", "effective_from"},
}
BOOLEAN_COLUMNS = {"active", "is_primary_technology", "is_buyer", "is_influencer", "is_budget_holder", "is_primary", "is_current"}
LIST_COLUMNS = {"skills", "tags"}
DATE_COLUMNS = {"effective_from", "effective_to"}
TEXT_ID_COLUMNS = {"id", "head_stakeholder_id", "division_id", "business_unit_id", "manager_employee_id", "manager_stakeholder_id", "stakeholder_id", "employee_id", "source_record_id"}


def _label(sheet: str, field: str) -> str:
    index = SHEETS[sheet].index(field)
    label = FRIENDLY_HEADERS[sheet][index]
    return f"{label} *" if field in REQUIRED_FIELDS[sheet] else label


def _normalize_header(value) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def _header_aliases(sheet: str) -> dict[str, str]:
    aliases = {}
    for field in SHEETS[sheet]:
        aliases[_normalize_header(field)] = field
        aliases[_normalize_header(_label(sheet, field))] = field
    return aliases


def template_bytes() -> bytes:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Instructions"
    guide.sheet_view.showGridLines = False
    guide.append(["Account import workbook", "Complete the seven data tabs, then preview the file in the app."])
    guide.append(["Step", "What to do"])
    for step, instruction in (
        ("1. Account", "Enter one client account name."),
        ("2. Pods and divisions", "Give each pod a name and a pod head contact ID. Add divisions under their pod names. Business units are optional."),
        ("3. Employees", "Enter your team members. Employee IDs stay the same wherever you reference them."),
        ("4. Client Contacts", "Enter client people. Each pod head has Organization role = Pod Head and no division, business unit, or manager."),
        ("5. Reporting lines", "Every other client contact needs a Division ID and Manager contact ID from this workbook."),
        ("6. Relationships", "Link a Client contact ID to an Employee ID. Leave this tab empty if ownership is not known yet."),
        ("7. Upload", "Choose Data administration > Excel import > Preview workbook. Import only after validation succeeds."),
    ):
        guide.append([step, instruction])
    guide.append(["How cells work", "Headers marked * are required when you add a row. Leave other cells blank if unknown."])
    guide.append(["IDs", "Use short text IDs such as pod-head-1 or employee-1. Reuse the exact text in linked ID columns; do not use a person's name as an ID."])
    guide.append(["Yes / No", "Choose Yes or No in dropdowns. Blank defaults to Yes for Active?, Influencer?, and Current?; the other flags default to No."])
    guide.append(["Dates and lists", "Use YYYY-MM-DD for dates. Separate multiple skills or tags with a semicolon (;)."])
    guide.append(["Examples", "See the Examples tab for a complete set of connected sample records. Do not copy its rows unless you want those sample names in your database."])
    guide.append(["Existing demo data", "The import needs an empty local database. Review demo data and complete both clear confirmations first; the app saves a backup."])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 110
    guide.freeze_panes = "A3"
    for cell in guide[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F59")
    for cell in guide[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176F9F")
    for row in guide.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[row[0].row].height = 32

    examples = workbook.create_sheet("Examples")
    examples.append(["Sheet", "Example record", "How it connects"])
    for row in (
        ("Account", "Example Client", "One account name."),
        ("Pods", "Banking | pod head contact ID: client-head", "The pod head is entered on Client Contacts."),
        ("Divisions", "division-1 | Banking | Lending", "Pod name must match Banking."),
        ("Business Units", "unit-1 | division-1 | Commercial Lending", "Division ID must match division-1."),
        ("Employees", "employee-1 | Alex Morgan", "Use employee-1 in Relationships."),
        ("Client Contacts", "client-head | Jordan Lee | Pod Head | Banking", "No manager or division for this row."),
        ("Client Contacts", "client-1 | Taylor Kim | Director | division-1 | manager: client-head", "This person reports to client-head."),
        ("Relationships", "relationship-1 | client-1 | employee-1 | Yes | 2026-01-01", "Links Taylor Kim to Alex Morgan."),
    ):
        examples.append(row)
    examples.freeze_panes = "A2"
    examples.column_dimensions["A"].width = 22
    examples.column_dimensions["B"].width = 82
    examples.column_dimensions["C"].width = 55
    for cell in examples[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176F9F")

    for sheet_name, fields in SHEETS.items():
        sheet = workbook.create_sheet(SHEET_TITLES[sheet_name])
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "176F9F" if sheet_name in {"Account", "Pods", "Divisions", "Stakeholders"} else "55A18A"
        sheet.append([_label(sheet_name, field) for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(fields)).column_letter}1"
        for field, cell in zip(fields, sheet[1]):
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176F9F" if field in REQUIRED_FIELDS[sheet_name] else "56758A")
            cell.alignment = Alignment(wrap_text=True)
            sheet.column_dimensions[cell.column_letter].width = min(34, max(18, len(str(cell.value)) + 3))
            hint = "Required" if field in REQUIRED_FIELDS[sheet_name] else "Optional"
            if field in BOOLEAN_COLUMNS:
                hint += ". Choose Yes or No."
            elif field in DATE_COLUMNS:
                hint += ". Use YYYY-MM-DD."
            elif field in LIST_COLUMNS:
                hint += ". Separate values with a semicolon (;)."
            elif field in TEXT_ID_COLUMNS:
                hint += ". Use a stable text ID and match it exactly in linked sheets."
            elif field == "team_type":
                hint += ". Choose Business or Technology."
            cell.comment = Comment(hint, "Account import")
            if field in BOOLEAN_COLUMNS or field in {"team_type", "relationship_strength"}:
                choices = "Yes,No" if field in BOOLEAN_COLUMNS else "Business,Technology" if field == "team_type" else "Strong,Medium,Developing,Unknown"
                validation = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
                validation.error = "Choose a value from the list."
                validation.errorTitle = "Invalid selection"
                validation.showErrorMessage = True
                sheet.add_data_validation(validation)
                validation.add(f"{cell.column_letter}2:{cell.column_letter}5001")
            if field in TEXT_ID_COLUMNS:
                for row in range(2, 102):
                    sheet.cell(row, cell.column).number_format = "@"
            if field in DATE_COLUMNS:
                for row in range(2, 102):
                    sheet.cell(row, cell.column).number_format = "yyyy-mm-dd"
        sheet.row_dimensions[1].height = 34
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_value(value, field: str, sheet: str, row: int):
    if value is None or value == "":
        return None
    location = f"{SHEET_TITLES[sheet]} row {row}, {_label(sheet, field).removesuffix(' *')}"
    if field in BOOLEAN_COLUMNS:
        if isinstance(value, bool):
            return value
        if str(value).strip().lower() in {"true", "yes", "1"}:
            return True
        if str(value).strip().lower() in {"false", "no", "0"}:
            return False
        raise ValueError(f"{location}: choose Yes or No")
    if field in LIST_COLUMNS:
        return [part.strip() for part in str(value).split(";") if part.strip()]
    if field in DATE_COLUMNS:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value).strip())
        except ValueError as error:
            raise ValueError(f"{location}: use YYYY-MM-DD") from error
    return value.strip() if isinstance(value, str) else value


def parse_workbook(content: bytes) -> PortableImportBatch:
    if not content or len(content) > 5 * 1024 * 1024:
        raise ValueError("Choose a nonempty .xlsx workbook no larger than 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("The file is not a readable .xlsx workbook") from error
    payload: dict = {"schema_version": 1}
    row_positions: dict[str, list[int]] = {}
    try:
        for sheet_name, fields in SHEETS.items():
            choices = [name for name in {sheet_name, SHEET_TITLES[sheet_name]} if name in workbook]
            if not choices:
                raise ValueError(f"Missing sheet: {SHEET_TITLES[sheet_name]}")
            if len(choices) > 1:
                raise ValueError(f"Keep only one {SHEET_TITLES[sheet_name]} sheet")
            sheet = workbook[choices[0]]
            if sheet.max_row and sheet.max_row > 5001:
                raise ValueError(f"{sheet.title} exceeds the 5,000 row limit")
            rows = sheet.iter_rows(values_only=True)
            actual = tuple(next(rows, ()))
            aliases = _header_aliases(sheet_name)
            field_columns = {}
            for index, header in enumerate(actual):
                if header is None or str(header).strip() == "":
                    continue
                field = aliases.get(_normalize_header(header))
                if not field:
                    raise ValueError(f"{sheet.title} has an unknown column: {header}")
                if field in field_columns:
                    raise ValueError(f"{sheet.title} has the same column twice: {_label(sheet_name, field)}")
                field_columns[field] = index
            missing_columns = REQUIRED_FIELDS[sheet_name] - field_columns.keys()
            if missing_columns:
                labels = ", ".join(_label(sheet_name, field) for field in fields if field in missing_columns)
                raise ValueError(f"{sheet.title} is missing required columns: {labels}")
            items = []
            key = {"BusinessUnits": "business_units", "Stakeholders": "stakeholders"}.get(sheet_name, sheet_name.lower())
            row_positions[key] = []
            for row_number, values in enumerate(rows, start=2):
                if not any(value is not None and value != "" for value in values):
                    continue
                if any(value is not None and value != "" for index, value in enumerate(values) if index not in field_columns.values()):
                    raise ValueError(f"{sheet.title} row {row_number}: data exists under a blank header")
                item = {}
                for field, index in field_columns.items():
                    parsed = _cell_value(values[index] if index < len(values) else None, field, sheet_name, row_number)
                    if parsed is not None:
                        item[field] = parsed
                missing_values = REQUIRED_FIELDS[sheet_name] - item.keys()
                if missing_values:
                    labels = ", ".join(_label(sheet_name, field) for field in fields if field in missing_values)
                    raise ValueError(f"{sheet.title} row {row_number}: fill required cells: {labels}")
                items.append(item)
                row_positions[key].append(row_number)
            payload[key] = items[0] if sheet_name == "Account" and len(items) == 1 else items
            if sheet_name == "Account" and len(items) != 1:
                raise ValueError("Account must have exactly one data row")
    finally:
        workbook.close()
    try:
        return PortableImportBatch.model_validate(payload)
    except ValidationError as error:
        detail = []
        for issue in error.errors()[:12]:
            parts = issue["loc"]
            key = parts[0] if parts else None
            source_sheet = next((name for name in SHEETS if {"BusinessUnits": "business_units"}.get(name, name.lower()) == key), None)
            if source_sheet and len(parts) > 1 and isinstance(parts[1], int) and parts[1] < len(row_positions.get(key, [])):
                location = f"{SHEET_TITLES[source_sheet]} row {row_positions[key][parts[1]]}"
                if len(parts) > 2 and parts[2] in SHEETS[source_sheet]:
                    location += f", {_label(source_sheet, parts[2])}"
            elif source_sheet and len(parts) > 1 and parts[1] in SHEETS[source_sheet]:
                location = f"{SHEET_TITLES[source_sheet]} row 2, {_label(source_sheet, parts[1])}"
            else:
                location = ".".join(str(part) for part in parts) or "Workbook"
            detail.append(f"{location}: {issue['msg']}")
        raise ValueError("; ".join(detail)) from error
